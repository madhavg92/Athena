from datetime import UTC, date, datetime, timedelta

from pptx import Presentation

from athena.app import build
from athena.core.model import ModelReply, StubModel
from athena.reports import wbr

MONDAY_8 = datetime(2026, 10, 5, 2, 30, tzinfo=UTC)  # Mon 08:00 IST


def make(cfg, db, tmp_path, model=None):
    root_cfg = cfg.model_copy(update={"root": tmp_path})
    (tmp_path / "context" / "templates").mkdir(parents=True)
    (tmp_path / "context" / "templates" / "messages.yaml").write_text("default: x\n")
    return build(cfg=root_cfg, db=db, clock=lambda: MONDAY_8, run_mode="fixture", model=model)


def test_last_week() -> None:
    assert wbr.last_week(MONDAY_8, "Asia/Kolkata") == (date(2026, 9, 28), date(2026, 10, 4))


def test_numbers_computed_in_code(cfg, db, tmp_path) -> None:
    app = make(cfg, db, tmp_path)
    n = wbr.compute(app, app.cfg.rules["R5"], "northwind_ortho")
    backlog = next(r for r in n.rows if r.metric == "backlog")
    week = [
        r.get("actual")
        for r in app.sources.read("supaboard.metrics", client="northwind_ortho", metric="backlog")
        if "2026-09-28" <= r.get("date") <= "2026-10-04"
    ]
    assert backlog.days == 7 and backlog.week_avg == round(sum(week) / 7, 1)
    assert (
        backlog.query_id == "wbr.northwind_ortho.backlog.2026-09-28" and backlog.off_target is True
    )
    assert not n.stale


def test_tick_builds_drafts_and_notifies_ba(cfg, db, tmp_path) -> None:
    from athena.core import scheduler

    app = make(cfg, db, tmp_path)
    out = scheduler.tick(app)
    r5 = next(r for r in out["runs"] if r["rule_id"] == "R5")
    assert r5["hits"] == 2 and r5["messages"] == {"sent": 2}
    path = tmp_path / "drafts" / "northwind_ortho-wbr-2026-09-28.pptx"
    prs = Presentation(str(path))
    assert len(prs.slides) == 3
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "wbr.northwind_ortho.backlog.2026-09-28" in notes
    to_ba = [m for m in app.gateway.deliverer.sent if m[0] == "ba.one@fixture.local"]
    assert len(to_ba) == 2 and "does not send it to the client" in to_ba[0][1]


def test_model_comments_with_invented_numbers_rejected(cfg, db, tmp_path) -> None:
    app = make(
        cfg,
        db,
        tmp_path,
        model=StubModel(script=[ModelReply(text="- Backlog fell 42% this week.")]),
    )
    n = wbr.compute(app, app.cfg.rules["R5"], "northwind_ortho")
    comments, *_ = wbr.write_comments(app, n)
    assert comments == wbr.template_comments(n)


def test_model_comments_with_real_numbers_kept(cfg, db, tmp_path) -> None:
    app = make(cfg, db, tmp_path)
    n = wbr.compute(app, app.cfg.rules["R5"], "northwind_ortho")
    backlog = next(r for r in n.rows if r.metric == "backlog")
    text = f"- Backlog averaged {backlog.week_avg:g} against {backlog.target:g}."
    app.model = StubModel(script=[ModelReply(text=text)])
    comments, *_ = wbr.write_comments(app, n)
    assert comments == [text.strip("- ")]


def test_stale_marked(cfg, db, tmp_path) -> None:
    rules = dict(cfg.rules)
    rules["R5"] = rules["R5"].model_copy(update={"data_max_age": timedelta(minutes=1)})
    app = make(cfg.model_copy(update={"rules": rules}), db, tmp_path)
    n = wbr.compute(app, app.cfg.rules["R5"], "bluefield_imaging")
    assert n.stale
    wbr.run_rule(app, app.cfg.rules["R5"])
    assert app.gateway.deliverer.sent[0][1].startswith("Data is not current")
