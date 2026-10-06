"""Weekly client report draft (R5).

Numbers are computed in code, each with a query ID in the slide notes. The model writes only the
comments, from the computed numbers. Any comment that has a number not in the computed set is
replaced by plain template comments. The draft goes to drafts/ and the BA is told. Athena never
sends it to the client.
"""

from __future__ import annotations

import json
import logging
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from pptx import Presentation
from pptx.util import Inches, Pt
from pydantic import BaseModel

from athena.app import App
from athena.connectors.base import NotConfigured
from athena.core.config import Rule
from athena.core.gateway import STALE_MARK, Action, Source

log = logging.getLogger(__name__)
DEFAULT_TEMPLATE = "context/templates/wbr.pptx"
COMMENT_SYSTEM = (
    "You write short comments for a weekly client report. Use only the numbers in the JSON. "
    "Do not compute new numbers. 3 to 5 bullet points, one line each. Never give patient details."
)


class MetricRow(BaseModel):
    metric: str
    meaning: str
    better: str
    week_avg: float | None
    target: float | None
    prior_avg: float | None
    change: float | None
    off_target: bool | None
    days: int
    query_id: str
    query: str


class WeeklyNumbers(BaseModel):
    client: str
    client_name: str
    week_start: date
    week_end: date
    rows: list[MetricRow]
    as_of: datetime | None
    stale: bool


def last_week(now: datetime, tz: str) -> tuple[date, date]:
    local = now.astimezone(ZoneInfo(tz)).date()
    this_monday = local - timedelta(days=local.weekday())
    start = this_monday - timedelta(days=7)
    return start, start + timedelta(days=6)


def _avg(values: list[float]) -> float | None:
    return round(sum(values) / len(values), 1) if values else None


def compute(app: App, rule: Rule, client: str) -> WeeklyNumbers:
    tz = app.cfg.work_hours_of(app.cfg.owner_map.clients[client].role("ba") or "")[1]
    now = app.clock()
    start, end = last_week(now, tz)
    prior_start, prior_end = start - timedelta(days=7), start - timedelta(days=1)
    records = app.sources.read(rule.source or "supaboard.metrics", client=client)
    as_of = min((r.as_of for r in records), default=None)
    rows = []
    for name, m in app.cfg.metrics.items():
        recs = [r for r in records if r.get("metric") == name]

        def values(a: date, b: date, recs=recs) -> list[float]:
            return [
                r.get("actual") for r in recs if a.isoformat() <= r.get("date") <= b.isoformat()
            ]

        week, prior = values(start, end), values(prior_start, prior_end)
        targets = [
            r.get("target") for r in recs if start.isoformat() <= r.get("date") <= end.isoformat()
        ]
        avg, prior_avg = _avg(week), _avg(prior)
        target = targets[-1] if targets else None
        off = None
        if avg is not None and target is not None:
            off = avg < target if m.better == "higher" else avg > target
        rows.append(
            MetricRow(
                metric=name,
                meaning=m.meaning,
                better=m.better,
                week_avg=avg,
                target=target,
                prior_avg=prior_avg,
                change=round(avg - prior_avg, 1)
                if avg is not None and prior_avg is not None
                else None,
                off_target=off,
                days=len(week),
                query_id=f"wbr.{client}.{name}.{start.isoformat()}",
                query=f"avg({rule.source}.actual) where client={client} and metric={name} and date in [{start}, {end}]",
            )
        )
    rule_max = rule.data_max_age
    stale = as_of is None or now - as_of > rule_max
    return WeeklyNumbers(
        client=client,
        client_name=app.cfg.owner_map.clients[client].name,
        week_start=start,
        week_end=end,
        rows=rows,
        as_of=as_of,
        stale=stale,
    )


def _fmt(value: float | None) -> str:
    return "-" if value is None else f"{value:g}"


def template_comments(n: WeeklyNumbers) -> list[str]:
    out = []
    for r in n.rows:
        if r.week_avg is None:
            out.append(f"{r.metric}: no data for the week.")
        elif r.off_target:
            out.append(
                f"{r.metric} averaged {_fmt(r.week_avg)} against a target of {_fmt(r.target)} ({r.better} is better)."
            )
        else:
            out.append(f"{r.metric} averaged {_fmt(r.week_avg)}, on target ({_fmt(r.target)}).")
    return out


def _numbers(n: WeeklyNumbers) -> set[str]:
    allowed = set()
    for r in n.rows:
        for v in (
            r.week_avg,
            r.target,
            r.prior_avg,
            r.change,
            abs(r.change) if r.change is not None else None,
            r.days,
        ):
            if v is not None:
                allowed |= {
                    f"{v:g}",
                    f"{float(v):.1f}",
                    str(int(v)) if float(v).is_integer() else f"{v:g}",
                }
    allowed |= {str(n.week_start.day), str(n.week_end.day), str(n.week_start.year), "7"}
    return allowed


def write_comments(app: App, n: WeeklyNumbers) -> tuple[list[str], int, int]:
    plain = template_comments(n)
    payload = {
        "client": n.client_name,
        "week": f"{n.week_start} to {n.week_end}",
        "metrics": [
            r.model_dump(
                include={
                    "metric",
                    "week_avg",
                    "target",
                    "prior_avg",
                    "change",
                    "off_target",
                    "better",
                }
            )
            for r in n.rows
        ],
        "template_hint": plain,
    }
    messages = [
        {"role": "system", "content": COMMENT_SYSTEM},
        {"role": "user", "content": json.dumps(payload, default=str)},
    ]
    try:
        reply = app.model.chat(messages, None)
    except Exception as exc:
        log.warning("comment writer failed", extra={"error": type(exc).__name__})
        return plain, 0, 0
    lines = [ln.strip(" -•\t") for ln in (reply.text or "").splitlines() if ln.strip(" -•\t")]
    found = set(re.findall(r"\d+(?:\.\d+)?", " ".join(lines)))
    if not lines or not found <= _numbers(n) or len(lines) > 8:
        return plain, reply.tokens_in, reply.tokens_out
    return lines, reply.tokens_in, reply.tokens_out


def make_default_template(path: Path) -> None:
    """A plain template until the BA team's standard one arrives (G7)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


def build_deck(app: App, rule: Rule, n: WeeklyNumbers, comments: list[str]) -> Path:
    template = app.cfg.root / (rule.builder.template if rule.builder else DEFAULT_TEMPLATE)
    if not template.exists():
        make_default_template(template)
    prs = Presentation(str(template))
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = f"Weekly business review: {n.client_name}"
    title.placeholders[
        1
    ].text = f"Week {n.week_start:%d %b} to {n.week_end:%d %b %Y}. DRAFT for BA review."

    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    table_slide.shapes.title.text = "Key metrics"
    headers = ["Metric", "Week average", "Target", "Prior week", "Change", "Status"]
    shape = table_slide.shapes.add_table(
        len(n.rows) + 1,
        len(headers),
        Inches(0.5),
        Inches(1.6),
        Inches(12.3),
        Inches(0.5) * (len(n.rows) + 1),
    )
    table = shape.table
    for c, h in enumerate(headers):
        table.cell(0, c).text = h
    for i, r in enumerate(n.rows, start=1):
        status = "-" if r.off_target is None else ("Off target" if r.off_target else "On target")
        for c, v in enumerate(
            [
                r.metric.replace("_", " "),
                _fmt(r.week_avg),
                _fmt(r.target),
                _fmt(r.prior_avg),
                _fmt(r.change),
                status,
            ]
        ):
            table.cell(i, c).text = v
            table.cell(i, c).text_frame.paragraphs[0].font.size = Pt(14)
    notes = ["Query IDs (numbers computed by Athena in code):"]
    notes += [f"{r.query_id}: {r.query} (days with data: {r.days})" for r in n.rows]
    if n.as_of:
        notes.append(f"Data as of {n.as_of:%Y-%m-%d %H:%M} UTC.")
    table_slide.notes_slide.notes_text_frame.text = "\n".join(notes)

    comment_slide = prs.slides.add_slide(prs.slide_layouts[1])
    comment_slide.shapes.title.text = "Comments"
    body = comment_slide.placeholders[1].text_frame
    lines = (
        [f"{STALE_MARK}: data as of {n.as_of:%Y-%m-%d %H:%M} UTC." if n.as_of else f"{STALE_MARK}."]
        if n.stale
        else []
    ) + comments
    body.text = lines[0]
    for line in lines[1:]:
        body.add_paragraph().text = line
    comment_slide.notes_slide.notes_text_frame.text = "Comments written by Athena from the numbers on the previous slide. The BA reviews and edits before sending."

    out_dir = app.cfg.root / "drafts"
    out_dir.mkdir(exist_ok=True)
    path = out_dir / f"{n.client}-wbr-{n.week_start.isoformat()}.pptx"
    prs.save(str(path))
    return path


def run_rule(app: App, rule: Rule):
    from athena.core.scheduler import RuleRun, _set_state

    result = RuleRun(rule_id=rule.id)
    clients = rule.clients or list(app.cfg.owner_map.clients)
    for client in clients:
        try:
            n = compute(app, rule, client)
        except NotConfigured as exc:
            result.stale, result.reason = True, f"source not available: {exc}"
            continue
        comments, t_in, t_out = write_comments(app, n)
        path = build_deck(app, rule, n, comments)
        rel = path.relative_to(app.cfg.root)
        ba = app.cfg.owner_map.clients[client].role(
            (rule.notify or "owner_map.client.ba").split(".")[-1]
        )
        text = (
            f"Draft weekly report for {n.client_name} (week of {n.week_start:%d %b}) is ready: {rel}. "
            "Please review and edit it before sending. Athena does not send it to the client."
        )
        if n.stale:
            text = f"{STALE_MARK}. " + text
        res = app.gateway.submit(
            Action(
                type="draft",
                rule_id=rule.id,
                actor="athena",
                recipients=[ba] if ba else [],
                clients=[client],
                text=text,
                payload={"path": str(rel)},
                delivery="now",
                mode=rule.mode,
                item_key=f"{rule.id}:{client}:{n.week_start}",
                sources=[
                    Source(
                        name=rule.source or "supaboard.metrics",
                        ref=str(rel),
                        as_of=n.as_of or app.clock(),
                    )
                ],
                tokens_in=t_in,
                tokens_out=t_out,
            )
        )
        status = "refused" if res.refused else next(iter(res.per_recipient.values()), "refused")
        result.messages[status] = result.messages.get(status, 0) + 1
        result.hits += 1
    _set_state(app, rule.id, last_run_at=app.clock())
    return result
